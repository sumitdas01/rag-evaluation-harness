import os
from pathlib import Path
from typing import List, Dict, Any
from langchain_core.documents import Document
from .cleaners import clean_text

class DocumentLoader:
    def __init__(self, documents_dir: str):
        self.documents_dir = Path(documents_dir)
        
    def load_all(self) -> List[Document]:
        """Scans the directory recursively and loads supported documents."""
        documents = []
        if not self.documents_dir.exists():
            print(f"Warning: Directory {self.documents_dir} does not exist.")
            return documents

        for root, _, files in os.walk(self.documents_dir):
            for file in files:
                file_path = Path(root) / file
                doc_id = str(file_path.relative_to(self.documents_dir))
                
                try:
                    if file.lower().endswith(".pdf"):
                        documents.extend(self._load_pdf(file_path, doc_id))
                    elif file.lower().endswith(".docx"):
                        documents.extend(self._load_docx(file_path, doc_id))
                    elif file.lower().endswith((".txt", ".md")):
                        documents.extend(self._load_text(file_path, doc_id))
                except Exception as e:
                    print(f"Error loading {file_path}: {e}")
                    
        return documents

    def _load_pdf(self, file_path: Path, doc_id: str) -> List[Document]:
        documents = []
        try:
            import fitz  # PyMuPDF
            with fitz.open(file_path) as pdf:
                for page_num in range(len(pdf)):
                    page = pdf[page_num]
                    text = page.get_text()
                    if text:
                        cleaned_text = clean_text(text)
                        metadata = {
                            "document_id": doc_id,
                            "source": str(file_path),
                            "file_name": file_path.name,
                            "file_type": "pdf",
                            "page": page_num + 1,
                        }
                        documents.append(Document(page_content=cleaned_text, metadata=metadata))
        except Exception as e:
            print(f"Failed to parse PDF {file_path}: {e}")
        return documents

    def _load_docx(self, file_path: Path, doc_id: str) -> List[Document]:
        documents = []
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            cleaned_text = clean_text("\n".join(full_text))
            if cleaned_text:
                metadata = {
                    "document_id": doc_id,
                    "source": str(file_path),
                    "file_name": file_path.name,
                    "file_type": "docx"
                }
                documents.append(Document(page_content=cleaned_text, metadata=metadata))
        except Exception as e:
            print(f"Failed to parse DOCX {file_path}: {e}")
        return documents

    def _load_text(self, file_path: Path, doc_id: str) -> List[Document]:
        documents = []
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            cleaned_text = clean_text(text)
            if cleaned_text:
                metadata = {
                    "document_id": doc_id,
                    "source": str(file_path),
                    "file_name": file_path.name,
                    "file_type": file_path.suffix[1:]
                }
                documents.append(Document(page_content=cleaned_text, metadata=metadata))
        except Exception as e:
            print(f"Failed to parse text file {file_path}: {e}")
        return documents
